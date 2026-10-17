"""
Semi-Automated Annotation via LLM Pre-Labeling

Sends flagged sentences to Claude Sonnet for contextual classification as
masculine, feminine, neutral, or both. Generates structured Excel workbooks
for human review and override. Processes sentences in batches of 15.

Author: Manya Reddy
Date:   2026-05-01
CS 72, Dartmouth College
"""

import pandas as pd
import anthropic
import json
import os
import time
import sys
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
BATCH_SIZE = 15
MODEL = "claude-sonnet-4-6"

client = anthropic.Anthropic(api_key=os.environ.get("ANTHROPIC_API_KEY", "YOUR_KEY_HERE"))

SYSTEM_PROMPT = """You are an expert annotator for a computational linguistics study on gendered language in job postings, based on Gaucher, Friesen & Kay (2011).

CONTEXT: We are building a dataset to train a BERT classifier to detect gendered language in job postings. We have a corpus of job postings from four industries (tech, finance, healthcare, education). Each posting has been segmented into sentences, and sentences have been automatically flagged if they contain words from established gendered word lists — specifically the masculine and feminine coded word stems from Gaucher et al. (2011), plus a small set of curated Word2Vec expansions.

However, keyword matching alone cannot distinguish between gendered and neutral uses of the same word. For example, "competitive" is masculine-coded when describing workplace culture ("competitive environment") but neutral when describing compensation ("competitive salary").

YOUR TASK: For each flagged sentence, validate whether the flagged word(s) are used in a genuinely gendered fashion or in a neutral/non-connotative way. If gendered, determine whether the usage is masculine-coded, feminine-coded, or both.

IMPORTANT: Gendered language is NOT limited to descriptions of the ideal candidate. Per Gaucher et al., the OVERALL wording of a job posting — including descriptions of company culture, work environment, and team values — affects applicants' sense of belongingness. A sentence like "We are an aggressive, fast-paced company" is masculine-coded because it signals a masculine workplace culture, even though it describes the organization.

GENDERED — label as masculine, feminine, or both when the word:
- Describes desired candidate traits/qualities ("assertive problem-solver")
- Describes company culture or work environment ("competitive environment", "supportive workplace")
- Describes team values or organizational identity ("we are an ambitious, driven team")
- Describes expected work style ("work independently", "collaborate with the team")

NEUTRAL — label as neutral ONLY when the word is used in a purely structural, factual, or non-connotative way:
- Compensation/benefits ("competitive salary", "generous benefits")
- Job duty descriptions ("responsible for managing budgets", "lead quarterly reviews")
- Legal/structural terms ("independent contractor", "decision pending approval")
- Non-adjective usage ("kind of work", "force majeure", "yield results")
- Factual market position ("leading provider of healthcare services")

Key examples:

MASCULINE:
- "competitive environment" = masculine (workplace culture)
- "We boast a competitive compensation package." = masculine (company description, from Gaucher et al.)
- "We are determined to deliver superior medical treatment tailored to each individual patient." = masculine (company description, from Gaucher et al.)
- "The superior candidate will have a self-confident attitude, decisive judgment" = masculine (desired traits, from Gaucher et al.)
- "Be self-reliant and able work independently" = masculine (expected work style, from Gaucher et al.)
- "In keeping with the hospital's hierarchy, administer medical interventions" = masculine (organizational framing, from Gaucher et al.)
- "aggressive growth strategy" = masculine (company values)

FEMININE:
- "supportive work environment" = feminine (workplace culture)
- "We support our employees with an excellent compensation package." = feminine (company description, from Gaucher et al.)
- "We are committed to provide top quality health care that considers the best needs of our patients." = feminine (company description, from Gaucher et al.)
- "A pleasant attitude, dependable judgment, and be attentive to details." = feminine (desired traits, from Gaucher et al.)
- "Connect with potential clients; being sensitive to their needs" = feminine (expected work style, from Gaucher et al.)
- "Collaborate with other care providers" = feminine (expected work style, from Gaucher et al.)
- "Caring and compassionate professional" = feminine (desired traits, from Gaucher et al.)

NEUTRAL — label as neutral ONLY for purely structural/factual usage:
- "competitive salary" = neutral (compensation amount, not culture)
- "responsible for managing budgets" = neutral (job duty, not a trait)
- "independent contractor" = neutral (legal/structural term)
- "decision pending approval" = neutral (process description)
- "kind of work" = neutral ("kind" as "type of", not the adjective)
- "leading provider of healthcare services" = neutral (factual market position)
- "Bilingualism is an asset." = neutral (no gendered words, from Gaucher et al.)
- "Bachelor of Nursing from recognized university." = neutral (factual requirement, from Gaucher et al.)

Labels:
- masculine: word(s) carry masculine-coded connotation (trait, culture, or work style)
- feminine: word(s) carry feminine-coded connotation (trait, culture, or work style)
- both: sentence contains masculine AND feminine coded words, both in gendered context
- neutral: word(s) are used in a purely structural/factual/non-connotative way

Respond with a JSON array. Each element must have:
- "index": the sentence index number provided
- "label": one of "masculine", "feminine", "both", "neutral"
- "reasoning": 1 brief sentence explaining why (especially for neutral — what non-gendered context is the word used in?)
"""


def build_batch_prompt(sentences):
    """Build a prompt for a batch of sentences."""
    lines = []
    for s in sentences:
        lines.append(
            f"[{s['idx']}] Sentence: \"{s['text']}\"\n"
            f"    Flagged words: {s['flagged_words']}\n"
            f"    (masculine: {s['masc_words']}, feminine: {s['fem_words']})"
        )
    return "Label each sentence below:\n\n" + "\n\n".join(lines)


def call_api(batch):
    """Send a batch to Claude and parse the response."""
    prompt = build_batch_prompt(batch)

    for attempt in range(3):
        try:
            response = client.messages.create(
                model=MODEL,
                max_tokens=3000,
                system=SYSTEM_PROMPT,
                messages=[{"role": "user", "content": prompt}],
            )
            text = response.content[0].text

            # Extract JSON from response
            start = text.find("[")
            end = text.rfind("]") + 1
            if start == -1 or end == 0:
                print(f"  WARNING: No JSON found in response, retrying...")
                continue

            results = json.loads(text[start:end])
            return results

        except json.JSONDecodeError:
            print(f"  WARNING: JSON parse error, retrying (attempt {attempt + 1}/3)...")
            time.sleep(2)
        except anthropic.RateLimitError:
            print(f"  Rate limited, waiting 30s...")
            time.sleep(30)
        except Exception as e:
            print(f"  ERROR: {e}, retrying (attempt {attempt + 1}/3)...")
            time.sleep(5)

    return None


def prelabel_sample(csv_path, label):
    """Pre-label all flagged sentences in a sample CSV."""
    print(f"\n{'=' * 60}")
    print(f"  Pre-labeling: {label}")
    print(f"{'=' * 60}")

    df = pd.read_csv(csv_path)
    flagged = df[df["flagged"] == True].copy()
    print(f"  Total sentences: {len(df):,}")
    print(f"  Flagged sentences to label: {len(flagged):,}")

    # Build batches
    batches = []
    current_batch = []
    for idx, row in flagged.iterrows():
        current_batch.append({
            "idx": idx,
            "text": row["sentence_text"],
            "flagged_words": row["all_flagged_words"],
            "masc_words": row["masculine_words"] if pd.notna(row["masculine_words"]) else "",
            "fem_words": row["feminine_words"] if pd.notna(row["feminine_words"]) else "",
        })
        if len(current_batch) >= BATCH_SIZE:
            batches.append(current_batch)
            current_batch = []
    if current_batch:
        batches.append(current_batch)

    print(f"  Batches: {len(batches)} (batch size: {BATCH_SIZE})")

    # Process batches
    all_results = {}
    for i, batch in enumerate(batches):
        print(f"  Batch {i + 1}/{len(batches)}...", end=" ", flush=True)
        results = call_api(batch)

        if results is None:
            print("FAILED — skipping batch")
            continue

        for r in results:
            all_results[r["index"]] = {
                "label": r["label"],
                "reasoning": r["reasoning"],
            }
        print(f"OK ({len(results)} labeled)")

        # Small delay between batches to avoid rate limits
        if i < len(batches) - 1:
            time.sleep(1)

    # Write results back to the dataframe
    df["suggested_label"] = ""
    df["reasoning"] = ""
    for idx, result in all_results.items():
        df.at[idx, "suggested_label"] = result["label"]
        df.at[idx, "reasoning"] = result["reasoning"]

    labeled_count = len(all_results)
    print(f"\n  Successfully labeled: {labeled_count}/{len(flagged)}")

    # Stats
    if all_results:
        labels = [r["label"] for r in all_results.values()]
        for l in ["masculine", "feminine", "both", "neutral"]:
            count = labels.count(l)
            print(f"    {l}: {count} ({count/len(labels)*100:.1f}%)")

    return df


def generate_workbook(df, output_path, label):
    """Generate an annotation workbook with pre-labels and reasoning."""
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    FLAGGED_FILL = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
    POSTING_SEPARATOR_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    MASC_FONT = Font(color="1565C0", bold=True)
    FEM_FONT = Font(color="C62828", bold=True)
    REASONING_FONT = Font(color="666666", italic=True, size=10)
    THIN_BORDER = Border(bottom=Side(style="thin", color="CCCCCC"))

    HEADERS = [
        "Job ID", "Title", "Sentence #", "Sentence Text", "Flagged",
        "Masculine Words", "Feminine Words",
        "Suggested Label", "Reasoning",
        "Your Label (final)", "Notes",
    ]
    COL_WIDTHS = {1: 14, 2: 30, 3: 10, 4: 80, 5: 8, 6: 25, 7: 25, 8: 16, 9: 50, 10: 16, 11: 30}

    wb = Workbook()
    wb.remove(wb.active)

    industries = sorted(df["industry_group"].unique())

    for industry in industries:
        ws = wb.create_sheet(title=industry.capitalize())
        industry_df = df[df["industry_group"] == industry]

        for col_num, width in COL_WIDTHS.items():
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(col_num)].width = width

        for col_idx, header in enumerate(HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:K1"

        row_num = 2
        postings = industry_df.groupby("job_id", sort=False)

        for job_id, posting_df in postings:
            posting_df = posting_df.sort_values("sentence_index")
            title = posting_df.iloc[0]["title"]

            sep_cell = ws.cell(row=row_num, column=1, value=f"--- Posting: {job_id} ---")
            sep_cell.font = Font(bold=True, size=10)
            for col_idx in range(1, len(HEADERS) + 1):
                ws.cell(row=row_num, column=col_idx).fill = POSTING_SEPARATOR_FILL
            ws.cell(row=row_num, column=2, value=title).font = Font(bold=True, size=10)
            flagged_count = posting_df["flagged"].sum()
            total_count = len(posting_df)
            ws.cell(row=row_num, column=5, value=f"{flagged_count}/{total_count}").font = Font(bold=True, size=10)
            row_num += 1

            for _, sent in posting_df.iterrows():
                is_flagged = sent["flagged"]

                ws.cell(row=row_num, column=1, value=sent["job_id"])
                ws.cell(row=row_num, column=2, value=sent["title"])
                ws.cell(row=row_num, column=3, value=sent["sentence_index"])

                text_cell = ws.cell(row=row_num, column=4, value=sent["sentence_text"])
                text_cell.alignment = Alignment(wrap_text=True, vertical="top")

                ws.cell(row=row_num, column=5, value="YES" if is_flagged else "")

                masc_cell = ws.cell(row=row_num, column=6, value=sent["masculine_words"] if pd.notna(sent["masculine_words"]) else "")
                if pd.notna(sent["masculine_words"]) and sent["masculine_words"]:
                    masc_cell.font = MASC_FONT

                fem_cell = ws.cell(row=row_num, column=7, value=sent["feminine_words"] if pd.notna(sent["feminine_words"]) else "")
                if pd.notna(sent["feminine_words"]) and sent["feminine_words"]:
                    fem_cell.font = FEM_FONT

                # Suggested label
                ws.cell(row=row_num, column=8, value=sent["suggested_label"] if pd.notna(sent["suggested_label"]) else "")

                # Reasoning
                reasoning_cell = ws.cell(row=row_num, column=9, value=sent["reasoning"] if pd.notna(sent["reasoning"]) else "")
                reasoning_cell.font = REASONING_FONT
                reasoning_cell.alignment = Alignment(wrap_text=True, vertical="top")

                # Your label (final)
                ws.cell(row=row_num, column=10, value="")

                # Notes
                ws.cell(row=row_num, column=11, value="")

                if is_flagged:
                    for col_idx in range(1, len(HEADERS) + 1):
                        ws.cell(row=row_num, column=col_idx).fill = FLAGGED_FILL

                for col_idx in range(1, len(HEADERS) + 1):
                    ws.cell(row=row_num, column=col_idx).border = THIN_BORDER

                row_num += 1
            row_num += 1

        dv = DataValidation(
            type="list",
            formula1='"masculine,feminine,both,neutral"',
            allow_blank=True,
        )
        dv.error = "Please select: masculine, feminine, both, or neutral"
        dv.errorTitle = "Invalid label"
        ws.add_data_validation(dv)
        dv.add(f"J2:J{row_num}")

    wb.save(output_path)
    print(f"  Saved: {output_path}")


# ──────────────────────────────────────────────
# Main
# ──────────────────────────────────────────────
print("Starting pre-labeling with Claude Sonnet...")

random_df = prelabel_sample(
    os.path.join(BASE_DIR, "sample_random.csv"), "Random Sample"
)
biased_df = prelabel_sample(
    os.path.join(BASE_DIR, "sample_biased.csv"), "Biased Sample"
)

print("\nGenerating annotation workbooks...")
generate_workbook(random_df, os.path.join(BASE_DIR, "annotation_random.xlsx"), "Random")
generate_workbook(biased_df, os.path.join(BASE_DIR, "annotation_biased.xlsx"), "Biased")

print("\nDone!")
