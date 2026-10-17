"""
Select Rewrite Candidates for Stage 2

Loads the enriched multiclass BERT model, recreates the train/test split
(seed=42), runs inference on the combined test set, ranks postings by
masculine sentence proportion (per Gaucher et al.), and exports the top 3
per industry (12 total) to Excel for LLM rewriting.

Author: Manya Reddy
Date:   2026-05-15
CS 72, Dartmouth College
"""

import os
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import torch
from torch.nn import functional as F
from transformers import BertTokenizer, BertForSequenceClassification
from sklearn.model_selection import train_test_split

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Model path: check both local and server structure
_multi_local = os.path.join(BASE_DIR, "bert_gendered", "output", "biased_multiclass", "checkpoint-best")
_multi_server = os.path.join(BASE_DIR, "output", "biased_multiclass", "checkpoint-best")
MULTI_MODEL_DIR = _multi_local if os.path.isdir(_multi_local) else _multi_server

# ──────────────────────────────────────────────
# Constants (same as train_bert.py)
# ──────────────────────────────────────────────
VALID_LABELS = {"masculine", "feminine", "both", "neutral"}
LABEL_TO_ID_MULTI = {"neutral": 0, "masculine": 1, "feminine": 2}


# ──────────────────────────────────────────────
# Data Loading (same logic as train_bert.py)
# ──────────────────────────────────────────────
def load_workbook_data(path, source_label):
    """Load all sentence rows from an annotation workbook."""
    print(f"  Loading {os.path.basename(path)}...")
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        industry = sheet_name.lower()

        for row in ws.iter_rows(min_row=2, values_only=True):
            cell_a = row[0]
            if cell_a is not None and str(cell_a).startswith("--- Posting:"):
                continue
            if cell_a is None or str(cell_a).strip() == "":
                continue
            sentence_text = row[3]
            if sentence_text is None or str(sentence_text).strip() == "":
                continue

            job_id = row[0]
            if isinstance(job_id, float):
                job_id = str(int(job_id))
            else:
                job_id = str(job_id)

            rows.append({
                "job_id": job_id,
                "title": row[1],
                "sentence_num": row[2],
                "sentence_text": str(sentence_text),
                "flagged": str(row[4]).strip().upper() == "YES" if row[4] else False,
                "masculine_words": row[5] or "",
                "feminine_words": row[6] or "",
                "suggested_label": str(row[7]).strip().lower() if row[7] else "",
                "your_label": str(row[9]).strip().lower() if row[9] else "",
                "industry": industry,
                "source": source_label,
            })

    return pd.DataFrame(rows)


def resolve_label(row):
    your = row["your_label"]
    if your in VALID_LABELS:
        return your
    suggested = row["suggested_label"]
    if suggested in VALID_LABELS:
        return suggested
    return "neutral"


def collapse_both(label):
    return "masculine" if label == "both" else label


def prepare_labels(df):
    df["resolved_label"] = df.apply(resolve_label, axis=1)
    df["resolved_label"] = df["resolved_label"].apply(collapse_both)
    df["multi_id"] = df["resolved_label"].map(LABEL_TO_ID_MULTI)
    return df


def split_by_posting(df, test_size=0.2, seed=42):
    posting_info = df.groupby("job_id").agg(
        industry=("industry", "first"),
        has_gendered=("multi_id", lambda x: int((x > 0).any())),
    ).reset_index()

    posting_info["strat_key"] = (
        posting_info["industry"] + "_" + posting_info["has_gendered"].astype(str)
    )

    strat_counts = posting_info["strat_key"].value_counts()
    rare_strata = strat_counts[strat_counts < 2].index
    if len(rare_strata) > 0:
        posting_info["strat_key"] = posting_info["industry"]

    train_ids, test_ids = train_test_split(
        posting_info["job_id"],
        test_size=test_size,
        random_state=seed,
        stratify=posting_info["strat_key"],
    )

    train_df = df[df["job_id"].isin(set(train_ids))].copy()
    test_df = df[df["job_id"].isin(set(test_ids))].copy()
    return train_df, test_df


# ──────────────────────────────────────────────
# Inference
# ──────────────────────────────────────────────
def run_inference_multiclass(model, tokenizer, texts, batch_size=32):
    """Run multiclass BERT inference, return full probability matrix."""
    model.eval()
    all_probs = []

    for i in range(0, len(texts), batch_size):
        batch_texts = texts[i:i + batch_size]
        encodings = tokenizer(
            batch_texts,
            padding="max_length",
            truncation=True,
            max_length=128,
            return_tensors="pt",
        )

        with torch.no_grad():
            outputs = model(**encodings)
            probs = F.softmax(outputs.logits, dim=1)
            all_probs.append(probs.cpu().numpy())

        if (i // batch_size) % 10 == 0:
            print(f"    Processed {min(i + batch_size, len(texts)):,}/{len(texts):,} sentences...", end="\r")

    print(f"    Processed {len(texts):,}/{len(texts):,} sentences.      ")
    return np.concatenate(all_probs, axis=0)


# ──────────────────────────────────────────────
# Posting Ranking (by MASCULINE proportion)
# ──────────────────────────────────────────────
def rank_postings(test_df):
    """Compute masculine proportion per posting and rank within industry."""
    # Aggregate at posting level
    posting_stats = test_df.groupby("job_id").agg(
        title=("title", "first"),
        industry=("industry", "first"),
        source=("source", "first"),
        total_sentences=("sentence_text", "count"),
        masculine_count=("predicted_masculine", "sum"),
        feminine_count=("predicted_feminine", "sum"),
        neutral_count=("predicted_neutral", "sum"),
        true_masculine=("resolved_label", lambda x: (x == "masculine").sum()),
        true_feminine=("resolved_label", lambda x: (x == "feminine").sum()),
    ).reset_index()

    posting_stats["masculine_proportion"] = (
        posting_stats["masculine_count"] / posting_stats["total_sentences"]
    )
    posting_stats["true_masculine_proportion"] = (
        posting_stats["true_masculine"] / posting_stats["total_sentences"]
    )

    # Rank within each industry by MASCULINE proportion
    posting_stats["rank"] = posting_stats.groupby("industry")["masculine_proportion"].rank(
        ascending=False, method="first"
    ).astype(int)

    posting_stats = posting_stats.sort_values(["industry", "rank"])
    return posting_stats


# ──────────────────────────────────────────────
# Export to Excel
# ──────────────────────────────────────────────
def export_candidates(posting_stats, test_df, output_path):
    """Export top 3 per industry to Excel with full sentence details."""
    top_postings = posting_stats[posting_stats["rank"] <= 3].copy()

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        # ── Sheet 1: Summary ──
        summary = top_postings[[
            "rank", "industry", "job_id", "title", "source",
            "total_sentences", "masculine_count", "masculine_proportion",
            "feminine_count", "true_masculine", "true_masculine_proportion"
        ]].copy()
        summary["masculine_proportion"] = (summary["masculine_proportion"] * 100).round(1)
        summary["true_masculine_proportion"] = (summary["true_masculine_proportion"] * 100).round(1)
        summary.columns = [
            "Rank", "Industry", "Job ID", "Title", "Source",
            "Total Sentences", "Masculine (BERT)", "% Masculine (BERT)",
            "Feminine (BERT)", "Masculine (Human)", "% Masculine (Human)"
        ]
        summary.to_excel(writer, sheet_name="Summary", index=False)

        # ── Sheet 2: All sentences from selected postings ──
        selected_ids = set(top_postings["job_id"])
        selected_sentences = test_df[test_df["job_id"].isin(selected_ids)].copy()
        selected_sentences = selected_sentences.sort_values(
            ["industry", "job_id", "sentence_num"]
        )

        sentences_export = selected_sentences[[
            "industry", "job_id", "title", "sentence_num", "sentence_text",
            "masculine_prob", "feminine_prob", "predicted_label",
            "resolved_label", "masculine_words", "feminine_words",
        ]].copy()
        sentences_export.columns = [
            "Industry", "Job ID", "Title", "Sentence #", "Sentence Text",
            "Masculine Prob", "Feminine Prob", "BERT Prediction (Multi)",
            "Human Label", "Masculine Words", "Feminine Words",
        ]
        sentences_export["Masculine Prob"] = sentences_export["Masculine Prob"].round(4)
        sentences_export["Feminine Prob"] = sentences_export["Feminine Prob"].round(4)
        sentences_export.to_excel(writer, sheet_name="All Sentences", index=False)

        # ── Sheet 3: Full posting rankings (all 120 postings) ──
        full_ranking = posting_stats[[
            "rank", "industry", "job_id", "title", "source",
            "total_sentences", "masculine_count", "masculine_proportion",
            "feminine_count", "true_masculine", "true_masculine_proportion"
        ]].copy()
        full_ranking["masculine_proportion"] = (full_ranking["masculine_proportion"] * 100).round(1)
        full_ranking["true_masculine_proportion"] = (full_ranking["true_masculine_proportion"] * 100).round(1)
        full_ranking.columns = [
            "Rank (in industry)", "Industry", "Job ID", "Title", "Source",
            "Total Sentences", "Masculine (BERT)", "% Masculine (BERT)",
            "Feminine (BERT)", "Masculine (Human)", "% Masculine (Human)"
        ]
        full_ranking.to_excel(writer, sheet_name="All Postings Ranked", index=False)

    print(f"\n  Saved: {output_path}")


# ──────────────────────────────────────────────
# Main
# ──────────────────────────────────────────────
def main():
    print("=" * 60)
    print("  SELECT REWRITE CANDIDATES (MASCULINE-FOCUSED)")
    print("=" * 60)

    # ── Load data ──
    print("\n  Loading annotation workbooks...")
    random_path = os.path.join(BASE_DIR, "annotation_random.xlsx")
    biased_path = os.path.join(BASE_DIR, "annotation_biased.xlsx")

    random_df = load_workbook_data(random_path, "random")
    random_df = prepare_labels(random_df)

    biased_df = load_workbook_data(biased_path, "biased")
    biased_df = prepare_labels(biased_df)

    # ── Split (same seed as training) ──
    print("\n  Splitting (posting-level, seed=42)...")
    _, random_test = split_by_posting(random_df)
    _, biased_test = split_by_posting(biased_df)

    # Combined test set
    test_df = pd.concat([random_test, biased_test], ignore_index=True)
    print(f"  Combined test set: {len(test_df):,} sentences from {test_df['job_id'].nunique()} postings")
    for ind in sorted(test_df["industry"].unique()):
        count = (test_df["industry"] == ind).sum()
        print(f"    {ind}: {count:,}")

    texts = test_df["sentence_text"].tolist()

    # ── Load and run multiclass model (for candidate selection) ──
    print(f"\n  Loading biased_multiclass model from:\n    {MULTI_MODEL_DIR}")
    multi_tokenizer = BertTokenizer.from_pretrained(MULTI_MODEL_DIR)
    multi_model = BertForSequenceClassification.from_pretrained(MULTI_MODEL_DIR)
    print(f"  Multiclass model loaded.")

    print("\n  Running multiclass inference...")
    multi_probs = run_inference_multiclass(multi_model, multi_tokenizer, texts)
    # Multiclass: 0=neutral, 1=masculine, 2=feminine
    test_df["neutral_prob"] = multi_probs[:, 0]
    test_df["masculine_prob"] = multi_probs[:, 1]
    test_df["feminine_prob"] = multi_probs[:, 2]
    test_df["predicted_class"] = np.argmax(multi_probs, axis=1)
    test_df["predicted_label"] = test_df["predicted_class"].map(
        {0: "neutral", 1: "masculine", 2: "feminine"}
    )
    test_df["predicted_masculine"] = (test_df["predicted_class"] == 1).astype(int)
    test_df["predicted_feminine"] = (test_df["predicted_class"] == 2).astype(int)
    test_df["predicted_neutral"] = (test_df["predicted_class"] == 0).astype(int)

    # ── Rank postings by MASCULINE proportion ──
    print("\n  Ranking postings by MASCULINE sentence density...")
    posting_stats = rank_postings(test_df)

    # Print top 3 per industry
    print("\n  Top 3 most MASCULINE-coded postings per industry:")
    print(f"  {'Industry':<12} {'Rank':<5} {'Job ID':<12} {'Masculine':<12} {'%':<8} {'Title'}")
    print(f"  {'-'*80}")
    top = posting_stats[posting_stats["rank"] <= 3]
    for _, row in top.iterrows():
        title = str(row["title"])[:35]
        print(f"  {row['industry']:<12} {row['rank']:<5} {row['job_id']:<12} "
              f"{int(row['masculine_count'])}/{int(row['total_sentences']):<7} "
              f"{row['masculine_proportion']*100:.1f}%    {title}")

    # ── Export ──
    output_path = os.path.join(BASE_DIR, "rewrite_candidates.xlsx")
    export_candidates(posting_stats, test_df, output_path)

    # ── Summary ──
    print("\n" + "=" * 60)
    print("  DONE")
    print("=" * 60)
    print(f"\n  Model: biased_multiclass (masculine / feminine / neutral)")
    print(f"  Selected 12 postings (3 per industry) by MASCULINE density")
    print(f"  Output: {output_path}")
    print(f"\n  Next steps:")
    print(f"    1. Open rewrite_candidates.xlsx")
    print(f"    2. Review the 'Summary' sheet — these are your Stage 2 candidates")
    print(f"    3. Check 'All Sentences' sheet — rewrite only the MASCULINE sentences")


if __name__ == "__main__":
    main()
